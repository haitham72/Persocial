"""Tier 1 triage outputs: triage.xlsx / triage.csv / triage.html.

Sheets: To review (verdict review), Hidden-suggested (verdict hide), Keep,
Audit (the audit sample, whatever its verdict), Summary (counts). Audit rows live
only in Audit so the verdict sheets sum cleanly; Summary counts everything.
"""

from __future__ import annotations

import csv
import html as _html
from pathlib import Path

COLUMNS = ["post", "comment_id", "username", "text", "language", "sentiment",
           "emotion", "signals", "verdict", "category", "severity",
           "confidence", "reason", "link"]

SUMMARY_SHEET = "Summary"


def instagram_link(post_code: str, ig_id: str) -> str:
    base = f"https://www.instagram.com/p/{post_code}/"
    if ig_id and ig_id.isdigit():
        return base + f"c/{ig_id}/"
    return base


def row_of(c: dict) -> dict:
    post = c.get("post_code") or ""
    ig_id = c.get("ig_id") or ""
    return {
        "post": post,
        "comment_id": ig_id,
        "username": c.get("username") or "",
        "text": c.get("text") or "",
        "language": c.get("script") or "",
        "sentiment": c.get("sentiment") or "",
        "emotion": c.get("emotion") or "",
        "signals": c.get("signals") or "",
        "verdict": c.get("verdict") or "",
        "category": c.get("category") or "",
        "severity": c.get("severity") if c.get("severity") is not None else "",
        "confidence": c.get("confidence") if c.get("confidence") is not None else "",
        "reason": c.get("reason") or "",
        "link": instagram_link(post, ig_id),
    }


def split_sheets(rows: list[dict]) -> dict[str, list[dict]]:
    out = {"To review": [], "Hidden-suggested": [], "Keep": [], "Audit": []}
    for c in rows:
        r = row_of(c)
        if (c.get("selection") or "") == "audit":
            out["Audit"].append(r)
        elif c.get("verdict") == "review":
            out["To review"].append(r)
        elif c.get("verdict") == "hide":
            out["Hidden-suggested"].append(r)
        else:
            out["Keep"].append(r)
    return out


def summary_rows(rows: list[dict], extra: dict | None = None) -> list[tuple[str, str]]:
    by_verdict: dict[str, int] = {}
    by_cat: dict[str, int] = {}
    for c in rows:
        by_verdict[c.get("verdict") or "?"] = by_verdict.get(c.get("verdict") or "?", 0) + 1
        by_cat[c.get("category") or "?"] = by_cat.get(c.get("category") or "?", 0) + 1
    posts = sorted({c.get("post_code") or "" for c in rows})
    out = [("comments", str(len(rows))), ("posts", str(len(posts)))]
    for v in ("hide", "review", "keep"):
        if v in by_verdict:
            out.append((f"verdict:{v}", str(by_verdict[v])))
    for cat in sorted(by_cat):
        out.append((f"category:{cat}", str(by_cat[cat])))
    for k, v in (extra or {}).items():
        out.append((k, str(v)))
    return out


def write_xlsx(path: Path, rows: list[dict], extra: dict | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    sheets = split_sheets(rows)
    for name, items in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in items:
            ws.append([r[c] for c in COLUMNS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    ws = wb.create_sheet(SUMMARY_SHEET)
    ws.append(["metric", "value"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for k, v in summary_rows(rows, extra):
        ws.append([k, v])
    wb.save(path)


def write_csv(path: Path, rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for c in rows:
            w.writerow(row_of(c))


def write_html(path: Path, rows: list[dict], extra: dict | None = None) -> None:
    template = (Path(__file__).parent / "templates" / "report.html").read_text(encoding="utf-8")
    data = [row_of(c) for c in rows]
    verdicts = sorted({r["verdict"] for r in data})
    cats = sorted({r["category"] for r in data})
    langs = sorted({r["language"] for r in data})
    summary = " · ".join(f"{_html.escape(k)}: {_html.escape(v)}"
                         for k, v in summary_rows(rows, extra))

    def opts(vals: list[str]) -> str:
        return "\n".join(f'<option value="{_html.escape(v)}">{_html.escape(v)}</option>'
                         for v in vals if v)

    body_rows = []
    for r in data:
        cells = []
        for c in COLUMNS:
            v = str(r[c])
            if c == "link":
                u = _html.escape(v, quote=True)
                cells.append(f'<td><a href="{u}">open</a></td>')
            else:
                cells.append(f"<td>{_html.escape(v)}</td>")
        body_rows.append("<tr>" + "".join(cells) + "</tr>")
    page = template.replace("%%SUMMARY%%", summary)
    page = page.replace("%%VERDICT_OPTS%%", opts(verdicts))
    page = page.replace("%%CATEGORY_OPTS%%", opts(cats))
    page = page.replace("%%LANGUAGE_OPTS%%", opts(langs))
    page = page.replace("%%ROWS%%", "\n".join(body_rows))
    page = page.replace("%%COLUMNS%%", ",".join(f'"{c}"' for c in COLUMNS))
    path.write_text(page, encoding="utf-8")
